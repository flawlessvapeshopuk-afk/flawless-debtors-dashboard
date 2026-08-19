#!/usr/bin/env python3
"""
Flawless UK Debtors Dashboard — processing script
Reads the latest .xlsx from the /data folder and writes index.html
"""

import json, re, sys, os, glob
import openpyxl

# ── Find the most recent xlsx in /data ──────────────────────────────────────
script_dir = os.path.dirname(os.path.abspath(__file__))
data_dir = os.path.join(script_dir, 'data')
xlsx_files = glob.glob(os.path.join(data_dir, '*.xlsx'))

if not xlsx_files:
    print("ERROR: No .xlsx files found in /data folder")
    sys.exit(1)

source_file = max(xlsx_files, key=os.path.getmtime)
print(f"Processing: {os.path.basename(source_file)}")

# ── Read workbook ────────────────────────────────────────────────────────────
try:
    wb = openpyxl.load_workbook(source_file, read_only=True, data_only=False)
except Exception as e:
    print(f"ERROR: Could not open file: {e}")
    sys.exit(1)

ws = wb.active
rows = list(ws.iter_rows(values_only=True))
print(f"Total rows in file: {len(rows)}")

# ── Find header row ──────────────────────────────────────────────────────────
header_idx = None
for i, row in enumerate(rows[:15]):
    row_lower = [str(c or '').lower().strip() for c in row]
    if 'customer id' in row_lower or 'customer name' in row_lower:
        header_idx = i
        print(f"Header row found at index {i}")
        break

if header_idx is None:
    print("ERROR: Could not find header row — dumping first 10 rows for diagnosis:")
    for i, row in enumerate(rows[:10]):
        print(f"  Row {i}: {[str(c or '')[:25] for c in row[:8]]}")
    sys.exit(1)

# ── Extract report date ──────────────────────────────────────────────────────
title = str(rows[0][0] or '') if rows else ''
m = re.search(r'upto\s+(\d+/\d+/\d+)', title, re.I)
report_date = m.group(1) if m else 'latest'
print(f"Report date: {report_date}")

headers = [str(c or '').strip() for c in rows[header_idx]]
print(f"Headers: {headers}")

# ── Find columns ─────────────────────────────────────────────────────────────
bank_col = next((i for i, h in enumerate(headers) if re.match(r'\d+/\d+/\d+', h)), None)
bank_date = headers[bank_col] if bank_col is not None else ''
print(f"Bank/payment column: index={bank_col}, date={bank_date}")

def ci(name):
    for i, h in enumerate(headers):
        if name.lower() in h.lower():
            return i
    return None

iId    = ci('customer id')
iName  = ci('customer name')
iRep   = ci('rep')
i30    = ci('0-30')
i31    = ci('31-60')
i61    = ci('61-90')
i91    = ci('91-120')
i121   = ci('121-')
iTotal = ci('total')

print(f"Column indices: id={iId} name={iName} rep={iRep} 0-30={i30} 31-60={i31} 61-90={i61} 91-120={i91} 121={i121} total={iTotal}")

if iName is None:
    print("ERROR: Cannot find 'Customer name' column")
    sys.exit(1)

# ── Parse numbers safely ─────────────────────────────────────────────────────
def num(row, i):
    if i is None or i >= len(row): return 0
    v = row[i]
    if v is None: return 0
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip()
    if s.startswith('='):
        s = s[1:]
        try: return sum(float(x) for x in s.replace('-', '+-').split('+') if x.strip())
        except: return 0
    try: return float(s.replace(',', ''))
    except: return 0

# ── Parse data rows ──────────────────────────────────────────────────────────
seen = set()
data = []
skipped = 0

for row_num, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
    try:
        name = str(row[iName] if iName is not None and iName < len(row) else '').strip()
        cid  = str(row[iId]   if iId   is not None and iId   < len(row) else '').strip()

        if not name: continue
        if any(t in name.lower() for t in ['grand total', 'total']): continue

        key = cid or name
        if key in seen:
            skipped += 1
            continue
        seen.add(key)

        b0   = num(row, i30)
        b31  = num(row, i31)
        b61  = num(row, i61)
        b91  = num(row, i91)
        b121 = num(row, i121)
        raw_total = num(row, iTotal)
        payment   = num(row, bank_col) if bank_col is not None else 0
        final_total = round(raw_total - payment, 2)

        if abs(final_total) < 0.01 and not any([b0, b31, b61, payment]):
            continue

        rep = str(row[iRep] if iRep is not None and iRep < len(row) else '').strip().title()

        data.append({
            "id": cid, "name": name, "rep": rep,
            "b0":   round(b0,   2),
            "b31":  round(b31,  2),
            "b61":  round(b61,  2),
            "b91":  round(b91,  2),
            "b121": round(b121, 2),
            "payment": round(payment, 2),
            "total": final_total
        })

    except Exception as e:
        print(f"  Warning: skipped row {row_num}: {e}")
        continue

print(f"Records parsed: {len(data)} (skipped {skipped} duplicates)")
print(f"Payments today: {sum(1 for r in data if r['payment'] > 0)}")

if len(data) == 0:
    print("ERROR: No data rows found — aborting so we don't publish a blank dashboard")
    sys.exit(1)

# ── HTML template ─────────────────────────────────────────────────────────────
HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Flawless UK &ndash; Debtors Dashboard</title>
<style>
*,*::before,*::after{box-sizing:border-box;margin:0;padding:0;}
:root{--bg:#f5f5f3;--surface:#fff;--surface1:#f9f9f8;--border:rgba(0,0,0,.10);--border-strong:rgba(0,0,0,.18);--text:#0b0b0b;--text-sec:#52514e;--text-muted:#898781;--accent:#2a78d6;--warn:#c98500;--danger:#d03b3b;--green:#2d7d3a;--green-bg:#EAF3DE;--radius:8px;}
@media(prefers-color-scheme:dark){:root{--bg:#1a1a19;--surface:#242422;--surface1:#2c2c2a;--border:rgba(255,255,255,.10);--border-strong:rgba(255,255,255,.18);--text:#f0efe8;--text-sec:#c3c2b7;--green-bg:#1a2e1a;}}
body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif;background:var(--bg);color:var(--text);font-size:14px;line-height:1.5;min-height:100vh;}
#wrap{padding:20px 24px 48px;max-width:1400px;margin:0 auto;}
.top-bar{display:flex;align-items:flex-start;justify-content:space-between;margin-bottom:18px;flex-wrap:wrap;gap:10px;}
.top-bar h1{font-size:20px;font-weight:700;letter-spacing:-.01em;}
.top-bar h1 span{color:var(--accent);}
.report-date{font-size:12px;color:var(--text-muted);margin-top:3px;}
.tag{display:inline-block;background:var(--surface1);border:0.5px solid var(--border-strong);border-radius:20px;padding:4px 12px;font-size:12px;font-weight:500;color:var(--text-sec);}
.kpi-row{display:grid;grid-template-columns:repeat(auto-fit,minmax(118px,1fr));gap:10px;margin-bottom:18px;}
.kpi{background:var(--surface1);border-radius:var(--radius);padding:.7rem 1rem;}
.kpi-label{font-size:11px;color:var(--text-muted);text-transform:uppercase;letter-spacing:.04em;margin-bottom:4px;}
.kpi-val{font-size:20px;font-weight:700;color:var(--text);}
.kpi-val.warn{color:var(--warn);}.kpi-val.danger{color:var(--danger);}.kpi-val.green{color:var(--green);}
.chart-row{display:grid;grid-template-columns:1.4fr 1fr;gap:12px;margin-bottom:18px;}
@media(max-width:680px){.chart-row{grid-template-columns:1fr;}}
.chart-card{background:var(--surface);border:0.5px solid var(--border);border-radius:var(--radius);padding:.85rem 1rem;}
.chart-title{font-size:11px;font-weight:600;color:var(--text-muted);text-transform:uppercase;letter-spacing:.05em;margin-bottom:10px;}
canvas{display:block;}
.rep-row{display:flex;align-items:center;gap:8px;padding:5px 0;border-bottom:0.5px solid var(--border);}
.rep-row:last-child{border-bottom:none;}
.rep-name{font-size:12px;color:var(--text-sec);width:140px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;flex-shrink:0;}
.bar-bg{flex:1;background:var(--surface1);border-radius:2px;height:7px;overflow:hidden;}
.bar-fill{height:100%;border-radius:2px;background:var(--accent);}
.rep-val{font-size:12px;font-weight:600;min-width:58px;text-align:right;color:var(--text);}
.controls{display:flex;gap:8px;flex-wrap:wrap;margin-bottom:10px;align-items:center;}
.controls input,.controls select{font-size:13px;padding:6px 10px;border:0.5px solid var(--border-strong);border-radius:var(--radius);background:var(--surface);color:var(--text);height:34px;outline:none;}
.controls input{width:200px;}.controls input:focus,.controls select:focus{border-color:var(--accent);}
.table-info{font-size:12px;color:var(--text-muted);margin-bottom:6px;}
.table-wrap{overflow-x:auto;border:0.5px solid var(--border);border-radius:var(--radius);}
table{width:100%;border-collapse:collapse;font-size:12px;}
th{background:var(--surface1);color:var(--text-muted);font-weight:600;font-size:11px;text-transform:uppercase;letter-spacing:.04em;padding:8px 10px;text-align:left;border-bottom:0.5px solid var(--border);white-space:nowrap;user-select:none;}
th.payment-col{color:var(--green);}
td{padding:7px 10px;border-bottom:0.5px solid var(--border);white-space:nowrap;color:var(--text);overflow:hidden;text-overflow:ellipsis;max-width:200px;}
tr:last-child td{border-bottom:none;}tr:hover td{background:var(--surface1);}
tr.has-payment td{background:rgba(46,125,50,0.04);}tr.has-payment:hover td{background:rgba(46,125,50,0.08);}
.num{text-align:right;}.payment-val{text-align:right;color:var(--green);font-weight:700;}
.badge{display:inline-block;font-size:11px;font-weight:600;padding:2px 7px;border-radius:4px;white-space:nowrap;}
.badge-ok{background:#EAF3DE;color:#3B6D11;}.badge-warn{background:#FAEEDA;color:#854F0B;}
.badge-danger{background:#FCEBEB;color:#A32D2D;}.badge-critical{background:#d03b3b;color:#fff;}
.num-warn{color:var(--warn);font-weight:600;}.num-danger{color:var(--danger);font-weight:600;}
.payment-tag{display:inline-flex;align-items:center;gap:4px;font-size:11px;font-weight:600;color:var(--green);background:var(--green-bg);border:1px solid rgba(45,125,58,0.25);border-radius:4px;padding:2px 7px;white-space:nowrap;}
.pagination{display:flex;gap:8px;align-items:center;margin-top:10px;font-size:12px;color:var(--text-sec);}
.pagination button{padding:4px 10px;border:0.5px solid var(--border-strong);border-radius:var(--radius);background:var(--surface);color:var(--text);font-size:12px;cursor:pointer;}
.pagination button:hover{background:var(--surface1);}.pagination button:disabled{opacity:.35;cursor:default;}
.footer{margin-top:32px;font-size:11px;color:var(--text-muted);text-align:center;}
</style>
</head>
<body>
<div id="wrap">
  <div class="top-bar">
    <div>
      <h1>Flawless UK &nbsp;<span>Debtors</span></h1>
      <div class="report-date">Aged debt as at {{REPORT_DATE}} &nbsp;&middot;&nbsp; Payments as at {{BANK_DATE}} &nbsp;&middot;&nbsp; {{CUSTOMER_COUNT}} customers</div>
    </div>
    <span class="tag">&#128202; Live report</span>
  </div>
  <div class="kpi-row" id="kpis"></div>
  <div class="chart-row">
    <div class="chart-card">
      <div class="chart-title">Balance by aging bucket</div>
      <canvas id="ageChart" height="190"></canvas>
    </div>
    <div class="chart-card">
      <div class="chart-title">Balance by sales rep</div>
      <div id="repBars"></div>
    </div>
  </div>
  <div class="controls">
    <input type="text" id="search" placeholder="Search customer, ID, rep&hellip;">
    <select id="repFilter"><option value="">All reps</option></select>
    <select id="bucketFilter">
      <option value="">All customers</option>
      <option value="current">Current only (0&ndash;30)</option>
      <option value="31">31&ndash;60 days overdue</option>
      <option value="61">61&ndash;90 days overdue</option>
      <option value="91">91&ndash;120 days overdue</option>
      <option value="121">121+ days overdue</option>
      <option value="overdue">Any overdue (31+)</option>
      <option value="paid">Payment received today</option>
    </select>
    <select id="sortField">
      <option value="total">Sort: balance &darr;</option>
      <option value="name">Sort: customer name</option>
      <option value="payment">Sort: payment &darr;</option>
      <option value="b31">Sort: 31&ndash;60 days &darr;</option>
      <option value="b61">Sort: 61&ndash;90 days &darr;</option>
      <option value="b91">Sort: 91+ days &darr;</option>
    </select>
  </div>
  <div class="table-info" id="tinfo"></div>
  <div class="table-wrap">
    <table>
      <thead><tr>
        <th style="min-width:175px">Customer</th>
        <th style="min-width:65px">ID</th>
        <th style="min-width:115px">Rep</th>
        <th class="num" style="min-width:78px">0&ndash;30</th>
        <th class="num" style="min-width:78px">31&ndash;60</th>
        <th class="num" style="min-width:78px">61&ndash;90</th>
        <th class="num" style="min-width:78px">91&ndash;120</th>
        <th class="num" style="min-width:78px">121+</th>
        <th class="num payment-col" style="min-width:105px">&#8595; Payment in</th>
        <th class="num" style="min-width:88px">Balance</th>
        <th style="min-width:100px">Status</th>
      </tr></thead>
      <tbody id="tbody"></tbody>
    </table>
  </div>
  <div class="pagination">
    <button id="prevBtn" disabled>&lsaquo; Prev</button>
    <span id="pageInfo"></span>
    <button id="nextBtn">Next &rsaquo;</button>
  </div>
  <div class="footer">Flawless UK &middot; Debtors Dashboard &middot; Aged debt {{REPORT_DATE}} &middot; Payments {{BANK_DATE}}</div>
</div>
<script>
var RAW={{DATA_JSON}};
var BANK_DATE='{{BANK_DATE_JS}}';
var page=0,PAGE=25;
function $(i){return document.getElementById(i);}
function pos(v){return v>0?v:0;}
function fmt(v){if(!v)return'&mdash;';var abs=Math.abs(v),s='&pound;'+abs.toLocaleString('en-GB',{minimumFractionDigits:0,maximumFractionDigits:0});return v<0?'-'+s:s;}
function fmtK(v){if(Math.abs(v)>=1000000)return'&pound;'+(v/1000000).toFixed(1)+'m';if(Math.abs(v)>=1000)return'&pound;'+Math.round(v/1000).toLocaleString('en-GB')+'k';return'&pound;'+Math.round(v).toLocaleString('en-GB');}
function drawBarChart(id,labels,values,colors){
  var canvas=$(id);if(!canvas)return;
  var dpr=window.devicePixelRatio||1,W=canvas.parentElement.getBoundingClientRect().width||500,H=190;
  canvas.width=W*dpr;canvas.height=H*dpr;canvas.style.width=W+'px';canvas.style.height=H+'px';
  var ctx=canvas.getContext('2d');ctx.scale(dpr,dpr);
  var pL=52,pR=12,pT=12,pB=48,cW=W-pL-pR,cH=H-pT-pB,max=Math.max.apply(null,values)||1;
  ctx.font='10px -apple-system,sans-serif';
  for(var i=0;i<=4;i++){
    var y=pT+cH-(cH*i/4),val=max*i/4;
    ctx.strokeStyle='rgba(150,150,140,0.2)';ctx.lineWidth=1;ctx.beginPath();ctx.moveTo(pL,y);ctx.lineTo(pL+cW,y);ctx.stroke();
    ctx.fillStyle='#898781';ctx.textAlign='right';
    ctx.fillText(val>=1000?('\u00a3'+Math.round(val/1000)+'k'):'\u00a30',pL-4,y+3);
  }
  var n=values.length,gap=8,bW=(cW-gap*(n+1))/n;
  for(var j=0;j<n;j++){
    var x=pL+gap*(j+1)+bW*j,bh=Math.max((values[j]/max)*cH,0),by=pT+cH-bh,r=Math.min(4,bh/2);
    ctx.beginPath();
    if(bh>0){ctx.moveTo(x+r,by);ctx.lineTo(x+bW-r,by);ctx.quadraticCurveTo(x+bW,by,x+bW,by+r);ctx.lineTo(x+bW,by+bh);ctx.lineTo(x,by+bh);ctx.lineTo(x,by+r);ctx.quadraticCurveTo(x,by,x+r,by);}
    else{ctx.rect(x,pT+cH-1,bW,1);}
    ctx.closePath();ctx.fillStyle=colors[j];ctx.fill();
    ctx.fillStyle='#898781';ctx.textAlign='center';ctx.font='10px -apple-system,sans-serif';
    var parts=labels[j].split(' ');
    ctx.fillText(parts[0],x+bW/2,pT+cH+14);if(parts[1])ctx.fillText(parts[1],x+bW/2,pT+cH+26);
  }
}
function buildRepFilter(){
  var el=$('repFilter');
  [...new Set(RAW.map(function(r){return r.rep;}).filter(Boolean))].sort().forEach(function(rep){var o=document.createElement('option');o.value=rep;o.textContent=rep;el.appendChild(o);});
}
function kpi(label,val,cls){return'<div class="kpi"><div class="kpi-label">'+label+'</div><div class="kpi-val '+cls+'">'+val+'</div></div>';}
function buildKPIs(){
  var total=RAW.reduce(function(s,r){return s+r.total;},0);
  var b0s=RAW.reduce(function(s,r){return s+pos(r.b0);},0);
  var b31s=RAW.reduce(function(s,r){return s+pos(r.b31);},0);
  var b61s=RAW.reduce(function(s,r){return s+pos(r.b61);},0);
  var b91s=RAW.reduce(function(s,r){return s+pos(r.b91);},0);
  var b121s=RAW.reduce(function(s,r){return s+pos(r.b121);},0);
  var payments=RAW.reduce(function(s,r){return s+r.payment;},0);
  var overdue=b31s+b61s+b91s+b121s,pct=total>0?Math.round(overdue/total*100):0;
  var crit=RAW.filter(function(r){return r.b91>0||r.b121>0;}).length;
  var paidCount=RAW.filter(function(r){return r.payment>0;}).length;
  $('kpis').innerHTML=
    kpi('Total balance',fmtK(total),'')
    +kpi('Current (0&ndash;30)',fmtK(b0s),'')
    +kpi('31&ndash;60 days',fmtK(b31s),b31s>0?'warn':'')
    +kpi('61&ndash;90 days',fmtK(b61s),b61s>0?'danger':'')
    +kpi('91+ days',fmtK(b91s+b121s),(b91s+b121s)>0?'danger':'')
    +kpi('Overdue %',pct+'%',pct>20?'danger':pct>10?'warn':'')
    +kpi(paidCount>0?'Paid today ('+paidCount+')':'Paid today',payments>0?fmtK(payments):'&mdash;',payments>0?'green':'')
    +kpi('91+ day accounts',''+crit,crit>0?'danger':'');
}
function buildChart(){
  drawBarChart('ageChart',
    ['0-30 days','31-60 days','61-90 days','91-120 days','121+ days'],
    [Math.round(RAW.reduce(function(s,r){return s+pos(r.b0);},0)),
     Math.round(RAW.reduce(function(s,r){return s+pos(r.b31);},0)),
     Math.round(RAW.reduce(function(s,r){return s+pos(r.b61);},0)),
     Math.round(RAW.reduce(function(s,r){return s+pos(r.b91);},0)),
     Math.round(RAW.reduce(function(s,r){return s+pos(r.b121);},0))],
    ['#2a78d6','#eda100','#eb6834','#e24b4a','#a32d2d']);
}
function buildRepBars(){
  var t={};RAW.forEach(function(r){var k=r.rep||'Unknown';t[k]=(t[k]||0)+r.total;});
  var sorted=Object.entries(t).filter(function(e){return e[1]>0;}).sort(function(a,b){return b[1]-a[1];}).slice(0,9);
  var max=sorted.length?sorted[0][1]:1;
  $('repBars').innerHTML=sorted.map(function(e){return'<div class="rep-row"><span class="rep-name" title="'+e[0]+'">'+e[0]+'</span><div class="bar-bg"><div class="bar-fill" style="width:'+Math.round(e[1]/max*100)+'%"></div></div><span class="rep-val">'+fmtK(e[1])+'</span></div>';}).join('');
}
function getFiltered(){
  var s=$('search').value.toLowerCase(),rep=$('repFilter').value,bucket=$('bucketFilter').value,sort=$('sortField').value;
  return RAW.filter(function(r){
    if(s&&r.name.toLowerCase().indexOf(s)<0&&r.id.toLowerCase().indexOf(s)<0&&r.rep.toLowerCase().indexOf(s)<0)return false;
    if(rep&&r.rep!==rep)return false;
    if(bucket==='current'&&r.b0<=0)return false;
    if(bucket==='31'&&r.b31<=0)return false;
    if(bucket==='61'&&r.b61<=0)return false;
    if(bucket==='91'&&r.b91<=0)return false;
    if(bucket==='121'&&r.b121<=0)return false;
    if(bucket==='overdue'&&(r.b31+r.b61+r.b91+r.b121)<=0)return false;
    if(bucket==='paid'&&r.payment<=0)return false;
    return true;
  }).sort(function(a,b){
    if(sort==='name')return a.name.localeCompare(b.name);
    if(sort==='payment')return b.payment-a.payment;
    if(sort==='b31')return b.b31-a.b31;
    if(sort==='b61')return b.b61-a.b61;
    if(sort==='b91')return(b.b91+b.b121)-(a.b91+a.b121);
    return b.total-a.total;
  });
}
function badge(r){
  if(r.b121>0)return'<span class="badge badge-critical">121+ days</span>';
  if(r.b91>0)return'<span class="badge badge-danger">91&ndash;120 days</span>';
  if(r.b61>0)return'<span class="badge badge-danger">61&ndash;90 days</span>';
  if(r.b31>0)return'<span class="badge badge-warn">31&ndash;60 days</span>';
  return'<span class="badge badge-ok">Current</span>';
}
function render(){
  var f=getFiltered(),start=page*PAGE,end=Math.min(start+PAGE,f.length);
  $('tinfo').innerHTML='Showing '+(f.length===0?0:start+1)+'&ndash;'+end+' of '+f.length+' customers';
  $('pageInfo').textContent='Page '+(page+1)+' of '+Math.max(1,Math.ceil(f.length/PAGE));
  $('prevBtn').disabled=page===0;$('nextBtn').disabled=end>=f.length;
  $('tbody').innerHTML=f.slice(start,end).map(function(r){
    var hp=r.payment>0;
    return'<tr class="'+(hp?'has-payment':'')+'">'
      +'<td title="'+r.name+'">'+r.name+'</td>'
      +'<td>'+(r.id||'&mdash;')+'</td>'
      +'<td title="'+r.rep+'">'+(r.rep||'&mdash;')+'</td>'
      +'<td class="num">'+(r.b0>0?fmt(r.b0):'&mdash;')+'</td>'
      +'<td class="num'+(r.b31>0?' num-warn':'')+'">'+(r.b31>0?fmt(r.b31):'&mdash;')+'</td>'
      +'<td class="num'+(r.b61>0?' num-danger':'')+'">'+(r.b61>0?fmt(r.b61):'&mdash;')+'</td>'
      +'<td class="num'+(r.b91>0?' num-danger':'')+'">'+(r.b91>0?fmt(r.b91):'&mdash;')+'</td>'
      +'<td class="num'+(r.b121>0?' num-danger':'')+'">'+(r.b121>0?fmt(r.b121):'&mdash;')+'</td>'
      +(hp?'<td class="payment-val">&minus;&pound;'+r.payment.toLocaleString('en-GB',{minimumFractionDigits:0,maximumFractionDigits:0})+'</td>':'<td class="num">&mdash;</td>')
      +'<td class="num" style="font-weight:700">'+fmt(r.total)+'</td>'
      +'<td>'+(hp?'<span class="payment-tag">&#10003; Paid '+BANK_DATE+'</span>':badge(r))+'</td>'
      +'</tr>';
  }).join('');
}
['search','repFilter','bucketFilter','sortField'].forEach(function(id){$(id).addEventListener(id==='search'?'input':'change',function(){page=0;render();});});
$('prevBtn').addEventListener('click',function(){page--;render();});
$('nextBtn').addEventListener('click',function(){page++;render();});
buildRepFilter();buildKPIs();buildChart();buildRepBars();render();
window.addEventListener('resize',function(){buildChart();});
</script>
</body>
</html>"""

# ── Inject data safely using placeholders ────────────────────────────────────
html = HTML_TEMPLATE
html = html.replace('{{REPORT_DATE}}',    report_date)
html = html.replace('{{BANK_DATE_JS}}',   bank_date)
html = html.replace('{{BANK_DATE}}',      bank_date)
html = html.replace('{{CUSTOMER_COUNT}}', str(len(data)))
html = html.replace('{{DATA_JSON}}',      json.dumps(data))

# ── Write output ─────────────────────────────────────────────────────────────
out_path = os.path.join(script_dir, 'index.html')
with open(out_path, 'w', encoding='utf-8') as f:
    f.write(html)

print(f"SUCCESS: Written index.html ({len(html):,} bytes)")
